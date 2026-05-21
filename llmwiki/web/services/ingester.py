"""Document upload, conversion, and wiki page creation."""

from __future__ import annotations

import re
import shutil
from datetime import date
from pathlib import Path

from llmwiki import REPO_ROOT

RAW_DIR = REPO_ROOT / "raw" / "company-docs"
WIKI_SOURCES = REPO_ROOT / "wiki" / "sources" / "company-docs"


def _slugify(text: str) -> str:
    text = re.sub(r'[\s]+', '-', text.strip())
    text = re.sub(r'[^\w\u4e00-\u9fff-]', '', text)
    return text.strip('-')[:50] or "unnamed"


def convert_docx(path: Path) -> str:
    """Convert .docx to markdown text."""
    from docx import Document
    doc = Document(str(path))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = para.style.name if para.style else ""
        if "Heading 1" in style:
            lines.append(f"# {text}")
        elif "Heading 2" in style:
            lines.append(f"## {text}")
        elif "Heading 3" in style:
            lines.append(f"### {text}")
        else:
            lines.append(text)
    # Extract tables
    for table in doc.tables:
        lines.append("")
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
    return "\n".join(lines)


def convert_pdf(path: Path) -> str:
    """Convert .pdf to markdown text."""
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    lines = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            lines.append(text)
    return "\n\n".join(lines)


def save_raw_file(filename: str, content: bytes) -> Path:
    """Save uploaded file to raw/company-docs/."""
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    dest = RAW_DIR / filename
    dest.write_bytes(content)
    return dest


def convert_doc(filepath: Path) -> str:
    """Convert legacy .doc to markdown. Tries multiple methods."""
    import subprocess
    import tempfile

    # Method 1: LibreOffice — most reliable for .doc with Chinese content
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            result = subprocess.run(
                [
                    "libreoffice", "--headless", "--convert-to", "docx",
                    "--outdir", tmpdir, str(filepath),
                ],
                capture_output=True, timeout=60,
            )
            if result.returncode == 0:
                # Find the converted .docx file
                converted = list(Path(tmpdir).glob("*.docx"))
                if converted:
                    docx_text = convert_docx(converted[0])
                    if docx_text and len(docx_text.strip()) > 50:
                        return docx_text
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass
    except Exception:
        pass

    # Method 2: Try antiword with UTF-8 mapping
    try:
        result = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", str(filepath)],
            capture_output=True, timeout=30,
        )
        if result.returncode == 0 and result.stdout:
            text = result.stdout.decode("utf-8", errors="replace")
            if len(text) > 100 and any('\u4e00' <= c <= '\u9fff' for c in text[:500]):
                return text
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    # Method 3: Try python-docx (works for some .doc files saved as docx-format)
    try:
        return convert_docx(filepath)
    except Exception:
        pass

    # Method 4: Extract text from OLE2 WordDocument stream
    try:
        import olefile
        ole = olefile.OleFileIO(str(filepath))
        if ole.exists("WordDocument"):
            for table_name in ["1Table", "0Table"]:
                if ole.exists(table_name):
                    raw = ole.openstream("WordDocument").read()
                    text_parts = []
                    i = 0
                    while i < len(raw) - 1:
                        char = raw[i:i+2]
                        try:
                            decoded = char.decode('utf-16-le')
                            if decoded.isprintable() or decoded in '\n\r\t':
                                text_parts.append(decoded)
                            else:
                                if text_parts and text_parts[-1] != '\n':
                                    text_parts.append('\n')
                        except UnicodeDecodeError:
                            if text_parts and text_parts[-1] != '\n':
                                text_parts.append('\n')
                        i += 2
                    text = ''.join(text_parts)
                    lines = [l.strip() for l in text.split('\n') if any('\u4e00' <= c <= '\u9fff' for c in l)]
                    if lines:
                        return '\n'.join(lines)
        ole.close()
    except ImportError:
        pass
    except Exception:
        pass

    raise ValueError(
        "无法转换 .doc 文件。建议：\n"
        "1. 用 Word 或 WPS 另存为 .docx 格式后重新上传\n"
        "2. 或联系管理员安装 LibreOffice"
    )


def convert_to_markdown(filepath: Path) -> str:
    """Convert uploaded file to markdown based on extension."""
    ext = filepath.suffix.lower()
    if ext == ".docx":
        return convert_docx(filepath)
    elif ext == ".doc":
        return convert_doc(filepath)
    elif ext == ".pdf":
        return convert_pdf(filepath)
    elif ext in (".md", ".txt"):
        return filepath.read_text(encoding="utf-8")
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


def create_source_page(filename: str, md_content: str, knowledge: dict | None = None) -> str:
    """Create a wiki source page from converted markdown content."""
    WIKI_SOURCES.mkdir(parents=True, exist_ok=True)

    stem = Path(filename).stem
    slug = _slugify(stem)
    today = date.today().isoformat()

    # Build structured page from GLM extraction
    summary = knowledge.get("summary", md_content[:200].replace("\n", " ").strip()) if knowledge else md_content[:200].replace("\n", " ").strip()
    claims = knowledge.get("key_claims", []) if knowledge else []
    entities = knowledge.get("entities", []) if knowledge else []
    concepts = knowledge.get("concepts", []) if knowledge else []

    claims_text = "\n".join(f"- {c}" for c in claims) if claims else "- （待 AI 提取）"
    connections = []
    for e in entities:
        connections.append(f"- [[{e['name']}]] — {e.get('description', '')}")
    for c in concepts:
        connections.append(f"- [[{c['name']}]] — {c.get('description', '')}")
    connections_text = "\n".join(connections) if connections else "- [[合同管理]]\n- [[采购管理]]"

    page_content = f"""---
title: "{stem}"
type: source
tags: [company-doc, internal-policy]
date: {today}
source_file: {filename}
project: company-docs
---

## Summary

{summary}

## Key Claims

{claims_text}

## Connections

{connections_text}

"""

    out_path = WIKI_SOURCES / f"{slug}.md"
    out_path.write_text(page_content + "\n" + md_content, encoding="utf-8")

    return f"sources/company-docs/{slug}.md"


def create_entity_concept_pages(knowledge: dict, source_slug: str) -> list[str]:
    """Create entity and concept pages from GLM extraction."""
    from llmwiki import REPO_ROOT as ROOT
    created = []
    today = date.today().isoformat()

    for item in knowledge.get("entities", []):
        name = item["name"]
        desc = item.get("description", "")
        etype = item.get("type", "organization")
        path = ROOT / "wiki" / "entities" / f"{name}.md"
        if path.exists():
            continue
        path.parent.mkdir(parents=True, exist_ok=True)
        content = f"""---
title: "{name}"
type: entity
tags: [{etype}]
sources: [{source_slug}]
last_updated: {today}
---

# {name}

{desc}

## Key Facts

- （从 {source_slug} 提取，待完善）

## Connections

- [[{source_slug}]]
"""
        path.write_text(content, encoding="utf-8")
        created.append(f"entities/{name}.md")

    for item in knowledge.get("concepts", []):
        name = item["name"]
        desc = item.get("description", "")
        path = ROOT / "wiki" / "concepts" / f"{name}.md"
        if path.exists():
            continue
        path.parent.mkdir(parents=True, exist_ok=True)
        content = f"""---
title: "{name}"
type: concept
tags: [management]
sources: [{source_slug}]
last_updated: {today}
---

# {name}

{desc}

## Key Facts

- （从 {source_slug} 提取，待完善）

## Connections

- [[{source_slug}]]
"""
        path.write_text(content, encoding="utf-8")
        created.append(f"concepts/{name}.md")

    return created


def _ensure_wikilink_pages(md_content: str, source_slug: str) -> list[str]:
    """Scan md_content for [[wikilinks]] and create stub entity pages for any
    that don't already exist in wiki/entities/ or wiki/concepts/."""
    from llmwiki import REPO_ROOT as ROOT

    WIKILINK_RE = re.compile(r"\[\[([^\]|]+)(?:\|[^\]]*)?\]\]")
    names = set(WIKILINK_RE.findall(md_content))
    if not names:
        return []

    today = date.today().isoformat()
    created = []

    for name in names:
        name = name.strip()
        if not name or "#" in name or "/" in name or "\\" in name or len(name) > 100:
            continue
        entity_path = ROOT / "wiki" / "entities" / f"{name}.md"
        concept_path = ROOT / "wiki" / "concepts" / f"{name}.md"
        if entity_path.exists() or concept_path.exists():
            continue
        # Create stub entity page
        entity_path.parent.mkdir(parents=True, exist_ok=True)
        content = f"""---
title: "{name}"
type: entity
tags: []
sources: [{source_slug}]
last_updated: {today}
---

# {name}

（自动生成的占位页面，待完善）

## Key Facts

- （从 {source_slug} 提取，待完善）

## Connections

- [[{source_slug}]]
"""
        entity_path.write_text(content, encoding="utf-8")
        created.append(f"entities/{name}.md")

    return created


def process_upload(filename: str, content: bytes) -> dict:
    """Full upload pipeline: save → convert → GLM extract → create wiki pages."""
    from llmwiki import REPO_ROOT as ROOT

    # 1. Save raw file
    raw_path = save_raw_file(filename, content)

    # 2. Convert to markdown
    try:
        md_text = convert_to_markdown(raw_path)
    except Exception as e:
        return {"status": "error", "message": f"转换失败: {e}"}

    # 3. Save markdown to raw/company-docs/
    md_filename = Path(filename).stem + ".md"
    md_path = RAW_DIR / md_filename
    today = date.today().isoformat()
    frontmatter = f"""---
title: "{Path(filename).stem}"
type: source
tags: [company-doc, internal-policy]
date: {today}
source_file: {filename}
project: company-docs
---

"""
    md_path.write_text(frontmatter + md_text, encoding="utf-8")

    # 4. GLM knowledge extraction
    knowledge = None
    entity_pages = []
    try:
        from llmwiki.web.services.glm_client import extract_knowledge
        knowledge = extract_knowledge(md_text)
    except Exception:
        pass  # Non-critical — source page still works without it

    # 5. Create wiki source page (with GLM-extracted metadata)
    wiki_path = create_source_page(filename, md_text, knowledge=knowledge)

    # 6. Create entity/concept pages from GLM extraction
    if knowledge:
        slug = Path(filename).stem
        entity_pages = create_entity_concept_pages(knowledge, slug)

    # 6b. Ensure all wikilinks in source page have corresponding pages
    source_page_path = ROOT / "wiki" / wiki_path
    if source_page_path.exists():
        source_content = source_page_path.read_text(encoding="utf-8")
        slug = Path(filename).stem
        wikilink_pages = _ensure_wikilink_pages(source_content, slug)
        entity_pages.extend(wikilink_pages)

    # 7. Append to log
    log_path = ROOT / "wiki" / "log.md"
    log_entry = f"\n## [{today}] ingest (web upload) | {Path(filename).stem}\n"
    if knowledge:
        log_entry += f"  - AI提取: {len(knowledge.get('entities', []))} 实体, {len(knowledge.get('concepts', []))} 概念\n"
    if entity_pages:
        log_entry += f"  - 新建页面: {', '.join(entity_pages)}\n"
    log_path.write_text(log_path.read_text(encoding="utf-8") + log_entry, encoding="utf-8")

    # Invalidate wiki search cache after writes
    from llmwiki.web.services.wiki_reader import invalidate_cache
    invalidate_cache()

    return {
        "status": "ok",
        "filename": filename,
        "wiki_path": wiki_path,
        "content_length": len(md_text),
        "entities_extracted": len(knowledge.get("entities", [])) if knowledge else 0,
        "concepts_extracted": len(knowledge.get("concepts", [])) if knowledge else 0,
        "new_pages": entity_pages,
    }
