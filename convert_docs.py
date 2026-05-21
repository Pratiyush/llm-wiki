"""Convert docx/pdf/doc files in D:/llm_wiki/raw/ to Markdown."""
import sys
from pathlib import Path

RAW_DIR = Path("D:/llm_wiki/raw")
OUT_DIR = Path("C:/Users/dayou/llm-wiki/raw/company-docs")

def convert_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue
        style = para.style.name if para.style else ""
        if "Heading 1" in style:
            lines.append(f"# {text}")
        elif "Heading 2" in style:
            lines.append(f"## {text}")
        elif "Heading 3" in style:
            lines.append(f"### {text}")
        elif "Heading 4" in style:
            lines.append(f"#### {text}")
        else:
            lines.append(text)
    # Extract tables
    for table in doc.tables:
        lines.append("")
        for i, row in enumerate(table.rows):
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if i == 0:
                lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
    return "\n".join(lines)

def convert_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    lines = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            lines.append(text)
        lines.append("")
    return "\n".join(lines)

def convert_doc(path: Path) -> str:
    # .doc (old format) — try reading as docx, fall back to raw text
    try:
        return convert_docx(path)
    except Exception:
        return path.read_text(encoding="utf-8", errors="ignore")

def slugify(name: str) -> str:
    import re
    name = re.sub(r'[^\w\u4e00-\u9fff-]', '-', name)
    name = re.sub(r'-+', '-', name).strip('-')
    return name

def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    converters = {
        ".docx": convert_docx,
        ".pdf": convert_pdf,
        ".doc": convert_doc,
    }

    converted = 0
    for f in sorted(RAW_DIR.iterdir()):
        if f.is_dir():
            continue
        ext = f.suffix.lower()
        if ext not in converters:
            print(f"  skip: {f.name} (unsupported format)")
            continue

        print(f"  converting: {f.name} ...", end=" ")
        try:
            md_content = converters[ext](f)
            slug = slugify(f.stem)
            out_path = OUT_DIR / f"{slug}.md"

            # Add frontmatter
            frontmatter = f"""---
title: "{f.stem}"
type: source
tags: [company-doc, internal-policy]
date: 2026-04-18
source_file: {f.name}
project: company-docs
---

"""
            out_path.write_text(frontmatter + md_content, encoding="utf-8")
            print(f"-> {out_path.name} ({len(md_content)} chars)")
            converted += 1
        except Exception as e:
            print(f"ERROR: {e}")

    print(f"\nConverted {converted} files to {OUT_DIR}")

if __name__ == "__main__":
    main()
